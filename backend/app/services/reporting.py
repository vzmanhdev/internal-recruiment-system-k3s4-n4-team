"""EP-09 reporting computed from database rows only (no synthesized numbers)."""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional

from sqlalchemy.orm import Session

from app.models.application import ApplicationLog, JobApplication
from app.models.hr import Department, JobTitle
from app.models.job import JobPosting
from app.models.master_data import RecruitmentCost
from app.models.offer_onboarding import JobOffer
from app.models.referral_eval import CandidateEvaluation, InternalReferral
from app.models.requisition import JobRequisition, RequisitionApprovalLog

STAGES = ["submitted", "screening", "interviewing", "offered", "hired"]
RANK = {s: i for i, s in enumerate(STAGES)}
RANK["onboarded"] = RANK["hired"]
STAGE_LABELS = {"submitted": "Nộp hồ sơ", "screening": "Sàng lọc", "interviewing": "Phỏng vấn", "offered": "Đề xuất offer", "hired": "Nhận việc"}


def _apps(db: Session, job_id=None, department_id=None, date_from: Optional[date] = None, date_to: Optional[date] = None) -> List[JobApplication]:
    q = db.query(JobApplication).join(JobPosting, JobPosting.id == JobApplication.job_posting_id)
    if job_id:
        q = q.filter(JobApplication.job_posting_id == job_id)
    if department_id:
        q = q.filter(JobPosting.department_id == department_id)
    if date_from:
        q = q.filter(JobApplication.created_at >= datetime.combine(date_from, datetime.min.time()))
    if date_to:
        q = q.filter(JobApplication.created_at <= datetime.combine(date_to, datetime.max.time()))
    return q.all()


def max_stage_reached(app: JobApplication) -> int:
    ranks = [RANK.get(app.status, -1)] + [RANK.get(log.to_status, -1) for log in app.logs]
    return max(0, *ranks)


def funnel(db: Session, **filters) -> dict:
    """S8-07: candidates that reached each stage, stage-to-stage conversion, biggest drop-off."""
    apps = _apps(db, **filters)
    reached = [max_stage_reached(a) for a in apps]
    rows = []
    for i, stage in enumerate(STAGES):
        count = sum(1 for r in reached if r >= i)
        prev = rows[-1]["count"] if rows else None
        rows.append({"stage": stage, "label": STAGE_LABELS[stage], "count": count,
                     "conversion_from_previous": round(count / prev * 100, 1) if prev else (100.0 if i == 0 and count else None)})
    drops = [(rows[i - 1]["count"] - rows[i]["count"], rows[i - 1]["stage"], rows[i]["stage"]) for i in range(1, len(rows)) if rows[i - 1]["count"]]
    worst = max(drops) if drops else None
    return {"total": len(apps), "stages": rows,
            "biggest_drop": {"from": worst[1], "to": worst[2], "lost": worst[0]} if worst and worst[0] > 0 else None}


def weekly_applications(db: Session, weeks: int = 8) -> List[dict]:
    today = date.today()
    start = today - timedelta(days=today.weekday()) - timedelta(weeks=weeks - 1)
    since = datetime.combine(start, datetime.min.time())
    buckets = {start + timedelta(weeks=i): 0 for i in range(weeks)}
    closed = {k: 0 for k in buckets}
    for (created,) in db.query(JobApplication.created_at).filter(JobApplication.created_at >= since).all():
        monday = created.date() - timedelta(days=created.weekday())
        if monday in buckets:
            buckets[monday] += 1
    for req in db.query(JobRequisition).filter(JobRequisition.status.in_(["filled", "cancelled"]), JobRequisition.updated_at >= since).all():
        monday = req.updated_at.date() - timedelta(days=req.updated_at.weekday())
        if monday in closed:
            closed[monday] += 1
    return [{"week_start": d.isoformat(), "applications": n, "closed_positions": closed[d]} for d, n in sorted(buckets.items())]


def monthly_closed_positions(db: Session, months: int = 6) -> List[dict]:
    today = date.today()
    keys = []
    y, m = today.year, today.month
    for _ in range(months):
        keys.append((y, m))
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    counts = {k: 0 for k in keys}
    for req in db.query(JobRequisition).filter(JobRequisition.status.in_(["filled", "cancelled"])).all():
        when = req.updated_at or req.created_at
        k = (when.year, when.month)
        if k in counts:
            counts[k] += 1
    return [{"month": f"{y:04d}-{m:02d}", "closed_positions": counts[(y, m)]} for (y, m) in sorted(keys)]


def _approval_date(db: Session, requisition_id: int) -> Optional[datetime]:
    row = db.query(RequisitionApprovalLog).filter(RequisitionApprovalLog.requisition_id == requisition_id,
                                                  RequisitionApprovalLog.action == "approved").order_by(RequisitionApprovalLog.id.desc()).first()
    return row.created_at if row else None


def hiring_metrics(db: Session, date_from: Optional[date] = None, date_to: Optional[date] = None) -> dict:
    """S8-08: time-to-hire (requisition approved → offer accepted) and cost-per-hire (manual costs + referral bonuses)."""
    titles = {t.id: t.name for t in db.query(JobTitle).all()}
    depts = {d.id: d.name for d in db.query(Department).all()}
    hires = []
    for offer in db.query(JobOffer).filter(JobOffer.status == "accepted", JobOffer.responded_at.isnot(None)).all():
        app = offer.application
        job = app.job_posting if app else None
        req = job.requisition if job else None
        if not req:
            continue
        if date_from and offer.responded_at.date() < date_from or date_to and offer.responded_at.date() > date_to:
            continue
        approved = _approval_date(db, req.id)
        days = (offer.responded_at - approved).days if approved else None
        hires.append({"requisition_id": req.id, "department": depts.get(req.department_id, "—"),
                      "job_title": req.job_title_name or titles.get(req.job_title_id) or req.title, "days": days, "application": app})
    costs = defaultdict(int)
    for c in db.query(RecruitmentCost).all():
        costs[c.requisition_id] += c.amount
    for ref in db.query(InternalReferral).all():
        app = db.query(JobApplication).filter(JobApplication.referral_code == f"REF-{ref.id}").first()
        if app and app.status in ("hired", "onboarded") and app.job_posting and app.job_posting.requisition_id:
            costs[app.job_posting.requisition_id] += ref.bonus_amount or 0

    def group(key):
        out = defaultdict(lambda: {"hires": 0, "days": [], "cost": 0, "reqs": set()})
        for h in hires:
            g = out[h[key]]
            g["hires"] += 1
            if h["days"] is not None:
                g["days"].append(h["days"])
            g["reqs"].add(h["requisition_id"])
        rows = []
        for name, g in out.items():
            cost = sum(costs.get(r, 0) for r in g["reqs"])
            rows.append({key: name, "hires": g["hires"],
                         "avg_time_to_hire_days": round(sum(g["days"]) / len(g["days"]), 1) if g["days"] else None,
                         "total_cost": cost, "cost_per_hire": round(cost / g["hires"]) if g["hires"] else None})
        return sorted(rows, key=lambda r: -r["hires"])

    all_days = [h["days"] for h in hires if h["days"] is not None]
    total_cost = sum(costs.get(r, 0) for r in {h["requisition_id"] for h in hires})
    return {
        "total_hires": len(hires),
        "avg_time_to_hire_days": round(sum(all_days) / len(all_days), 1) if all_days else None,
        "cost_per_hire": round(total_cost / len(hires)) if hires else None,
        "by_department": group("department"),
        "by_job_title": group("job_title"),
    }


def source_effectiveness(db: Session, date_from: Optional[date] = None, date_to: Optional[date] = None) -> List[dict]:
    """S8-09: per source — volume, screening pass rate, hires and a quality score (avg. evaluation), not only volume."""
    groups: Dict[str, List[JobApplication]] = defaultdict(list)
    for app in _apps(db, date_from=date_from, date_to=date_to):
        groups[app.source or "Không rõ"].append(app)
    rows = []
    for source, apps in groups.items():
        passed = sum(1 for a in apps if max_stage_reached(a) >= RANK["interviewing"])
        hired = sum(1 for a in apps if a.status in ("hired", "onboarded"))
        scores = [e.weighted_score if e.weighted_score is not None else (e.overall_score or 0) * 10
                  for e in db.query(CandidateEvaluation).filter(CandidateEvaluation.application_id.in_([a.id for a in apps])).all()]
        rows.append({"source": source, "applications": len(apps), "passed_screening": passed,
                     "screening_pass_rate": round(passed / len(apps) * 100, 1), "hired": hired,
                     "hire_rate": round(hired / len(apps) * 100, 1),
                     "avg_evaluation_score": round(sum(scores) / len(scores), 1) if scores else None})
    # quality first (hire rate, then evaluation score), volume last
    rows.sort(key=lambda r: (-r["hire_rate"], -(r["avg_evaluation_score"] or 0), -r["applications"]))
    if rows:
        rows[0]["best_quality"] = True
    return rows


def to_xlsx(sheets: Dict[str, List[dict]]) -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    col_name_map = {
        "department": "Phòng Ban",
        "job_title": "Chức Danh Tuyển Dụng",
        "hires": "Số Lượng Tuyển Được",
        "avg_time_to_hire_days": "Thời Gian Tuyển TB (Ngày)",
        "total_cost": "Tổng Chi Phí (VND)",
        "cost_per_hire": "Chi Phí / Tuyển Dụng (VND)",
        "source": "Nguồn Tuyển Dụng",
        "applications": "Số Lượng Ứng Tuyển",
        "passed_screening": "Qua Vòng Hồ Sơ",
        "screening_pass_rate": "Tỷ Lệ Qua Hồ Sơ (%)",
        "hired": "Trúng Tuyển",
        "hire_rate": "Tỷ Lệ Trúng Tuyển (%)",
        "avg_evaluation_score": "Điểm Đánh Giá TB",
        "best_quality": "Nguồn Tốt Nhất",
    }

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.views.sheetView[0].showGridLines = True
        if rows:
            raw_headers = list(rows[0].keys())
            display_headers = [col_name_map.get(h, h) for h in raw_headers]
            ws.append(display_headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
            ws.row_dimensions[1].height = 28

            for row_idx, r in enumerate(rows, start=2):
                row_vals = []
                for h in raw_headers:
                    val = r.get(h)
                    if isinstance(val, bool):
                        val = "✓ Đạt" if val else ""
                    elif val is None:
                        val = "—"
                    row_vals.append(val)
                ws.append(row_vals)
                for col_idx, cell in enumerate(ws[row_idx], start=1):
                    cell.font = regular_font
                    cell.border = thin_border
                    raw_h = raw_headers[col_idx - 1]
                    if raw_h in ("total_cost", "cost_per_hire") and isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[row_idx].height = 22

            # Auto fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        else:
            ws.append(["Chưa có dữ liệu phát sinh trong kỳ"])
            ws.column_dimensions["A"].width = 35

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
