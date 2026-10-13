"""Feature completion (EP-04..EP-09): real API calls + independent DB reads. SMTP disabled → e-mails are queued
and, when the outbox runs, recorded as 'skipped' (never 'sent')."""
import base64
import io
import json
import zipfile
from datetime import date, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.main import app
from app.db.session import Base
from app.db.seed_rbac import ROLES_DATA, PERMISSIONS_DATA, DEFAULT_GRANTED_MATRIX
from app.models.user import User
from app.models.rbac import Role, UserRole
from app.models.audit import Permission, RolePermission
from app.models.hr import Department, JobTitle, RecruiterAssignment, SharedCategory
from app.models.requisition import JobRequisition, RequisitionApprovalLog
from app.models.job import JobPosting
from app.models.application import JobApplication
from app.models.email_automation import EmailLog
from app.models.interview import InterviewSchedule
from app.models.notification import Notification
from app.models.master_data import TalentPoolEntry, SlaRule
from app.models.referral_eval import CandidateEvaluation
from app.core.security import hash_password
from app.services.email_dispatch import process_email_outbox
from app.services.scheduled_jobs import expire_postings, run_sla_check
from tests.conftest import session_token

engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


client = TestClient(app)
PDF = b"%PDF-1.4\n1 0 obj\nstream\nBT (Nguyen Van An) Tj ET\nBT (an.nguyen@example.com) Tj ET\nBT (0912345678) Tj ET\nendstream\n%%EOF"
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def docx_bytes(text: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", f"<w:document><w:body><w:p><w:t>{text}</w:t></w:p></w:body></w:document>")
    return buf.getvalue()


@pytest.fixture(autouse=True)
def w():
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    roles = {r["code"]: Role(code=r["code"], name=r["name"]) for r in ROLES_DATA}
    db.add_all(roles.values())
    for p in PERMISSIONS_DATA:
        db.add(Permission(code=p["code"], name=p["name"], module=p["module"]))
    for role_code, granted in DEFAULT_GRANTED_MATRIX.items():
        for p in PERMISSIONS_DATA:
            db.add(RolePermission(role_code=role_code, permission_code=p["code"], is_granted=p["code"] in granted))
    db.commit()
    users = {}
    for key, role in [("hrm", "hr_manager"), ("rec", "recruiter"), ("rec2", "recruiter"), ("hm", "hiring_manager"),
                      ("int", "interviewer"), ("int2", "interviewer"), ("emp", "employee")]:
        u = User(email=f"{key}@example.com", full_name=key.upper(), hashed_password=hash_password("Password123"), role=role,
                 is_active=True, account_status="active", security_version=1)
        db.add(u); db.flush()
        db.add(UserRole(user_id=u.id, role_id=roles[role].id))
        users[key] = u
    dept = Department(code="ENG", name="Kỹ thuật"); db.add(dept); db.commit()
    title = JobTitle(name="Kỹ sư Backend", department_id=dept.id, level="Senior", min_salary=20_000_000, max_salary=30_000_000)
    db.add(title); db.commit()
    req = JobRequisition(code="REQ-F-1", title="Backend", department_id=dept.id, job_title_id=title.id, job_title_name=title.name,
                         quantity=2, status="approved", current_approval_step="completed", created_by_id=users["hm"].id,
                         min_salary=20_000_000, max_salary=30_000_000, created_at=datetime.utcnow() - timedelta(days=20))
    db.add(req); db.commit()
    db.add(RequisitionApprovalLog(requisition_id=req.id, approver_id=users["hrm"].id, step_name="HR", action="approved",
                                  created_at=datetime.utcnow() - timedelta(days=15)))
    db.add(RecruiterAssignment(requisition_id=req.id, recruiter_id=users["rec"].id, assignment_type="primary"))
    job = JobPosting(slug="backend", title="Kỹ sư Backend", requisition_id=req.id, department_id=dept.id, department_name=dept.name,
                     description="Phát triển API với FastAPI", status="published", created_by_id=users["hrm"].id,
                     min_years_experience=3, required_skills_json=json.dumps(["python", "fastapi"]))
    db.add(job); db.commit()
    out = {"h": {k: {"Authorization": f"Bearer {session_token(db, u)}"} for k, u in users.items()},
           "ids": {k: u.id for k, u in users.items()}, "dept": dept.id, "req": req.id, "job": job.id, "title": title.id}
    db.close()
    yield out
    Base.metadata.drop_all(bind=engine)


def upload(content=PDF, kind="cv", name="cv.pdf", headers=None):
    return client.post("/api/v1/files", data={"kind": kind}, files={"file": (name, content, "application/pdf")}, headers=headers or {})


def apply(w, email="an@example.com", cv=True, headers=None, **extra):
    body = {"job_posting_id": w["job"], "full_name": "Nguyễn Văn An", "email": email, "phone": "0912345678",
            "years_of_experience": 4, "current_position": "Python FastAPI developer"}
    if cv:
        body["cv_url"] = upload().json()["url"]
    body.update(extra)
    return client.post("/api/v1/applications", json=body, headers=headers or {})


def db_get(model, **f):
    db = TestingSessionLocal()
    try:
        return db.query(model).filter_by(**f).first()
    finally:
        db.close()


# ------------------------------------------------------------------ files (S2-03, S4-03, S4-04)
def test_upload_validates_content_size_and_access(w):
    ok = upload()
    assert ok.status_code == 201 and ok.json()["content_type"] == "application/pdf"
    assert ok.json()["parsed"]["email"] == "an.nguyen@example.com"           # S4-04 parsed for confirmation
    assert upload(content=b"MZ\x90\x00 not a pdf").status_code == 415           # renamed executable
    assert upload(content=b"%PDF-" + b"0" * (10 * 1024 * 1024 + 1)).status_code == 413
    assert upload(content=docx_bytes("Tran Thi B  b@example.com"), name="cv.docx").status_code == 201
    assert upload(content=PNG, kind="avatar").status_code == 401               # avatar needs an account
    url = ok.json()["url"]
    assert client.get(url).status_code == 404                                  # anonymous cannot download
    assert client.get(url, headers=w["h"]["emp"]).status_code == 404           # unrelated user cannot download
    assert client.get(url, headers=w["h"]["hrm"]).status_code == 200


def test_cv_visible_to_scoped_recruiter_only(w):
    cv = upload().json()["url"]
    client.post("/api/v1/applications", json={"job_posting_id": w["job"], "full_name": "X Y", "email": "x@example.com",
                                              "phone": "0900000000", "cv_url": cv})
    assert client.get(cv, headers=w["h"]["rec"]).status_code == 200
    assert client.get(cv, headers=w["h"]["rec2"]).status_code == 404


def test_profile_avatar_must_be_valid_image(w):
    bad = "data:image/svg+xml;base64," + base64.b64encode(b"<svg onload=alert(1)>").decode()
    assert client.put("/api/v1/auth/me", headers=w["h"]["emp"], json={"avatar_url": bad}).status_code == 415
    good = "data:image/png;base64," + base64.b64encode(PNG).decode()
    assert client.put("/api/v1/auth/me", headers=w["h"]["emp"], json={"avatar_url": good}).status_code == 200
    assert client.put("/api/v1/auth/me", headers=w["h"]["emp"], json={"avatar_url": "https://evil.example/x.png"}).status_code == 422


# ------------------------------------------------------------------ portal (S4-01..S4-06, S5-09)
def test_application_requires_uploaded_cv_and_gets_tracking_code_and_email(w):
    assert apply(w, cv=False, cv_url="https://elsewhere.example/cv.pdf").status_code == 422
    resp = apply(w)
    assert resp.status_code == 201
    code = resp.json()["tracking_code"]
    assert code.startswith("APP-") and "hr_notes" not in resp.json()
    stored = db_get(JobApplication, tracking_code=code)
    assert stored.match_score == 100 and stored.meets_requirements is True     # S5-09
    log = db_get(EmailLog, template_code="CONFIRM_APPLICATION", related_id=str(stored.id))
    assert log.status == "queued" and code in log.body_html                    # outbox row in the same transaction
    db = TestingSessionLocal()
    assert process_email_outbox(db) >= 1
    db.close()
    assert db_get(EmailLog, id=log.id).status == "skipped"                     # SMTP off → never reported as sent


def test_public_status_lookup_reveals_only_general_status(w):
    code = apply(w).json()["tracking_code"]
    ok = client.post("/api/v1/public/applications/status", json={"tracking_code": code, "email": "an@example.com"})
    assert ok.status_code == 200 and ok.json()["status"] == "Đã nhận" and set(ok.json()) == {"tracking_code", "job_title", "status", "applied_date"}
    assert client.post("/api/v1/public/applications/status", json={"tracking_code": code, "email": "other@example.com"}).status_code == 404


def test_withdraw_through_email_link_requires_confirmation_and_is_single_use(w):
    resp = apply(w)
    app_id = resp.json()["id"]
    db = TestingSessionLocal()
    from app.core.security import hash_token
    # the raw token only exists in the e-mail: recover it from the queued body for the test
    body = db.query(EmailLog).filter(EmailLog.related_id == str(app_id)).first().body_html
    db.close()
    token = body.split("token=")[1].split('"')[0]
    preview = client.post("/api/v1/public/applications/withdraw", json={"token": token})
    assert preview.json()["confirm_required"] is True and db_get(JobApplication, id=app_id).status == "submitted"
    done = client.post("/api/v1/public/applications/withdraw", json={"token": token, "confirm": True})
    assert done.status_code == 200 and db_get(JobApplication, id=app_id).status == "withdrawn"
    assert client.post("/api/v1/public/applications/withdraw", json={"token": token, "confirm": True}).status_code == 404
    assert db_get(Notification, type="application_withdrawn").recipient_id == w["ids"]["rec"]


def test_anti_spam_honeypot_and_hourly_ip_limit(w):
    assert apply(w, website="http://spam").status_code == 422
    for i in range(5):
        assert apply(w, email=f"c{i}@example.com", cv=False).status_code == 201
    assert apply(w, email="c9@example.com", cv=False).status_code == 429


def test_public_portal_hides_expired_and_searches_description(w):
    db = TestingSessionLocal()
    db.add(JobPosting(slug="old", title="Cũ", requisition_id=w["req"], department_id=w["dept"], department_name="Kỹ thuật",
                      description="hết hạn", status="published", created_by_id=w["ids"]["hrm"], deadline="2020-01-01"))
    db.commit(); db.close()
    titles = [j["title"] for j in client.get("/api/v1/jobs?search=FastAPI").json()]
    assert titles == ["Kỹ sư Backend"]
    assert client.get("/api/v1/jobs/old").json()["is_expired"] is True
    db = TestingSessionLocal()
    assert expire_postings(db) == 1
    db.close()
    assert db_get(JobPosting, slug="old").status == "closed"


# ------------------------------------------------------------------ candidate pipeline (S4-09, S5-01..S5-08, S6-01, S7-09)
def test_internal_notes_tags_profile_and_candidate_view(w):
    app_id = apply(w, headers=w["h"]["emp"]).json()["id"]
    client.post(f"/api/v1/applications/{app_id}/notes", headers=w["h"]["rec"], json={"content": "Ứng viên mạnh về API"})
    client.put(f"/api/v1/applications/{app_id}/tags", headers=w["h"]["rec"], json={"tags": ["python", "senior"], "rating": 4})
    prof = client.get(f"/api/v1/applications/{app_id}/profile", headers=w["h"]["rec"]).json()
    assert prof["notes"][0]["author_name"] == "REC" and prof["application"]["rating"] == 4
    assert any(t["type"] == "applied" for t in prof["timeline"])
    mine = client.get(f"/api/v1/applications/{app_id}", headers=w["h"]["emp"]).json()
    assert "hr_notes" not in mine and "rating" not in mine and mine["public_status"] == "Đã nhận"
    assert client.get(f"/api/v1/applications/{app_id}/notes", headers=w["h"]["emp"]).status_code == 403


def test_duplicate_detection_merge_and_undo(w):
    a = apply(w, email="dup@example.com", cv=False).json()["id"]
    db = TestingSessionLocal()
    other_job = JobPosting(slug="be2", title="Backend 2", requisition_id=w["req"], department_id=w["dept"], department_name="Kỹ thuật",
                           description="d", status="published", created_by_id=w["ids"]["hrm"])
    db.add(other_job); db.commit(); job2 = other_job.id; db.close()
    b = client.post("/api/v1/applications", json={"job_posting_id": job2, "full_name": "An khác", "email": "dup2@example.com",
                                                  "phone": "0912345678"}).json()["id"]
    dups = client.get(f"/api/v1/applications/{a}/duplicates", headers=w["h"]["hrm"]).json()["duplicates"]
    assert [d["id"] for d in dups] == [b]
    assert client.post(f"/api/v1/applications/{a}/merge/{b}", headers=w["h"]["hrm"]).status_code == 200
    assert db_get(JobApplication, id=b).person_key == "dup@example.com"
    assert len(client.get(f"/api/v1/applications/{a}/profile", headers=w["h"]["hrm"]).json()["previous_applications"]) == 1
    assert client.post(f"/api/v1/applications/{b}/unmerge", headers=w["h"]["hrm"]).status_code == 200
    assert db_get(JobApplication, id=b).person_key == "dup2@example.com"


def test_stage_rules_rejection_catalog_skip_guard_and_talent_pool(w):
    app_id = apply(w, cv=False).json()["id"]
    assert client.put(f"/api/v1/applications/{app_id}/status", headers=w["h"]["rec"], json={"status": "offered"}).status_code == 409
    db = TestingSessionLocal()
    db.add(SharedCategory(category_type="rejection_reason", name="Không đủ kinh nghiệm")); db.commit(); db.close()
    bad = client.put(f"/api/v1/applications/{app_id}/status", headers=w["h"]["rec"], json={"status": "rejected", "rejection_reason": "tuỳ ý"})
    assert bad.status_code == 422
    ok = client.put(f"/api/v1/applications/{app_id}/status", headers=w["h"]["rec"], json={
        "status": "rejected", "rejection_reason": "Không đủ kinh nghiệm", "add_to_talent_pool": True, "talent_pool_reason": "Tiềm năng 2027"})
    assert ok.status_code == 200 and db_get(TalentPoolEntry, application_id=app_id) is not None
    sugg = client.get(f"/api/v1/talent-pool/suggestions?requisition_id={w['req']}", headers=w["h"]["hrm"]).json()
    assert sugg and sugg[0]["application_id"] == app_id


def test_export_excludes_internal_notes_and_guards_formulas(w):
    apply(w, cv=False, full_name="=HYPERLINK(\"http://x\")")
    resp = client.get(f"/api/v1/applications/export?job_id={w['job']}", headers=w["h"]["hrm"])
    assert resp.status_code == 200 and "JOB" in resp.headers["content-disposition"] and date.today().isoformat() in resp.headers["content-disposition"]
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
    header = [c.value for c in ws[1]]
    assert "Ghi chú" not in " ".join(header) and str(ws.cell(row=2, column=2).value).startswith("'=")


def test_bulk_rejection_letters_preview_and_no_duplicates(w):
    ids = [apply(w, email=f"r{i}@example.com", cv=False).json()["id"] for i in range(2)]
    preview = client.post("/api/v1/applications/bulk-email", headers=w["h"]["rec"], json={"application_ids": ids, "preview_only": True}).json()
    assert len(preview["items"]) == 2 and "r0@example.com" in preview["items"][0]["email"]
    first = client.post("/api/v1/applications/bulk-email", headers=w["h"]["rec"], json={"application_ids": ids}).json()
    second = client.post("/api/v1/applications/bulk-email", headers=w["h"]["rec"], json={"application_ids": ids}).json()
    assert first["queued"] == 2 and second["queued"] == 0 and second["skipped_duplicates"] == 2


def test_manual_candidate_and_saved_filters(w):
    resp = client.post("/api/v1/applications/manual", headers=w["h"]["rec"], json={
        "job_posting_id": w["job"], "full_name": "Sourced", "email": "an@example.com", "phone": "0912345678"})
    assert resp.status_code == 201 and resp.json()["source"] == "Recruiter nhập tay"
    assert client.post("/api/v1/applications/manual", headers=w["h"]["rec2"], json={
        "job_posting_id": w["job"], "full_name": "S2", "email": "s2@example.com", "phone": "1"}).status_code == 403
    client.post("/api/v1/applications/saved-filters", headers=w["h"]["rec"], json={"name": "Python", "filters": {"tag": "python"}})
    assert client.get("/api/v1/applications/saved-filters", headers=w["h"]["rec"]).json()[0]["filters"] == {"tag": "python"}


# ------------------------------------------------------------------ interviews (S6-02..S6-09, S7-01..S7-03)
def _when(hours):
    return (datetime.now() + timedelta(hours=hours)).replace(microsecond=0).isoformat()


def _framework(w):
    fw = client.post("/api/v1/hr/competency-frameworks", headers=w["h"]["hrm"], json={
        "code": "CF-BE", "name": "Backend", "applied_job_titles": ["Kỹ sư Backend"],
        "criteria": [{"code": "A", "name": "Thiết kế", "weight": 60}, {"code": "B", "name": "Code", "weight": 40}]}).json()
    q = client.post("/api/v1/hr/interview-questions", headers=w["h"]["hrm"], json={
        "code": "Q1", "content": "Thiết kế hệ thống?", "criteria_id": fw["criteria"][0]["id"]}).json()
    return fw, q


def test_invitation_link_confirmation_room_conflict_and_reschedule(w):
    app_id = apply(w, cv=False).json()["id"]
    first = client.post("/api/v1/interviews", headers=w["h"]["rec"], json={
        "application_id": app_id, "scheduled_time": _when(48), "location": "Phòng 302", "interviewer_ids": [w["ids"]["int"]]})
    assert first.status_code == 201
    other = apply(w, email="o@example.com", cv=False).json()["id"]
    clash = client.post("/api/v1/interviews", headers=w["h"]["rec"], json={
        "application_id": other, "scheduled_time": _when(48), "location": "phòng 302", "interviewer_ids": [w["ids"]["int2"]]})
    assert clash.status_code == 409 and clash.json()["error"]["code"] == "INTERVIEW_CONFLICT"
    overridden = client.post("/api/v1/interviews", headers=w["h"]["rec"], json={
        "application_id": other, "scheduled_time": _when(48), "location": "phòng 302", "interviewer_ids": [w["ids"]["int2"]],
        "override_conflict_reason": "Phòng lớn chia đôi"})
    assert overridden.status_code == 201

    log = db_get(EmailLog, template_code="INVITE_INTERVIEW", related_id=str(app_id))
    token = log.body_html.split("token=")[1].split('"')[0]
    assert client.get(f"/api/v1/public/interviews?token={token}").json()["location"] == "Phòng 302"
    assert client.post("/api/v1/public/interviews/respond", json={"token": token, "response": "confirmed"}).status_code == 200
    assert db_get(InterviewSchedule, id=first.json()["id"]).candidate_response == "confirmed"

    iid = first.json()["id"]
    resched = client.put(f"/api/v1/interviews/{iid}", headers=w["h"]["rec"], json={
        "scheduled_time": _when(72), "location": "Phòng 401", "interviewer_ids": [w["ids"]["int"], w["ids"]["hm"]]})
    assert resched.status_code == 200 and len(resched.json()["interviewers"]) == 2
    assert client.post(f"/api/v1/interviews/{iid}/cancel", headers=w["h"]["rec"], json={"reason": ""}).status_code == 422


def test_bulk_scheduling_preview_then_confirm(w):
    ids = [apply(w, email=f"b{i}@example.com", cv=False).json()["id"] for i in range(3)]
    body = {"application_ids": ids, "start_time": _when(24), "slot_minutes": 30, "interviewer_ids": [w["ids"]["int"]], "location": "Phòng 101"}
    preview = client.post("/api/v1/interviews/bulk", headers=w["h"]["rec"], json=body).json()
    starts = [datetime.fromisoformat(s["start"]) for s in preview["slots"]]
    assert preview["preview"] and starts[1] - starts[0] == timedelta(minutes=30)
    assert db_get(InterviewSchedule, application_id=ids[0]) is None
    done = client.post("/api/v1/interviews/bulk", headers=w["h"]["rec"], json={**body, "confirm": True}).json()
    assert len(done["created"]) == 3
    db = TestingSessionLocal()
    assert db.query(EmailLog).filter(EmailLog.template_code == "INVITE_INTERVIEW").count() == 3
    db.close()


def test_competency_scorecard_weighted_score_snapshot_reopen_and_compare(w):
    fw, q = _framework(w)
    app_id = apply(w, cv=False).json()["id"]
    iid = client.post("/api/v1/interviews", headers=w["h"]["rec"], json={
        "application_id": app_id, "scheduled_time": _when(30), "interviewer_ids": [w["ids"]["int"]]}).json()["id"]
    tpl = client.get(f"/api/v1/interviews/{iid}/scorecard-template", headers=w["h"]["int"]).json()
    assert [c["weight"] for c in tpl["criteria"]] == [60, 40]
    assert client.get(f"/api/v1/interviews/{iid}/scorecard-template", headers=w["h"]["int2"]).status_code == 403
    sugg = client.get(f"/api/v1/interviews/{iid}/question-suggestions", headers=w["h"]["int"]).json()
    assert sugg[0]["code"] == "Q1" and sugg[0]["already_asked"] is False
    client.post(f"/api/v1/interviews/{iid}/asked-questions", headers=w["h"]["int"], json={"question_id": q["id"]})
    assert client.get(f"/api/v1/interviews/{iid}/question-suggestions", headers=w["h"]["int"]).json()[0]["already_asked"] is True

    crit = tpl["criteria"]
    partial = client.post("/api/v1/referrals/evaluations", headers=w["h"]["int"], json={
        "application_id": app_id, "interview_id": iid, "recommendation": "hire", "ratings": [{"criterion_id": crit[0]["id"], "score": 5}]})
    assert partial.status_code == 422
    ev = client.post("/api/v1/referrals/evaluations", headers=w["h"]["int"], json={
        "application_id": app_id, "interview_id": iid, "recommendation": "hire",
        "ratings": [{"criterion_id": crit[0]["id"], "score": 5}, {"criterion_id": crit[1]["id"], "score": 3}]})
    assert ev.status_code == 201 and ev.json()["weighted_score"] == 84.0   # (5*60 + 3*40) / 100 / 5 * 100
    # framework changes do not alter submitted evaluations (S6-07 AC4)
    client.put(f"/api/v1/hr/competency-frameworks/{fw['id']}", headers=w["h"]["hrm"], json={
        "code": "CF-BE", "name": "Backend", "applied_job_titles": ["Kỹ sư Backend"],
        "criteria": [{"id": crit[0]["id"], "code": "A", "name": "Thiết kế mới", "weight": 50}, {"id": crit[1]["id"], "code": "B", "name": "Code", "weight": 50}]})
    assert json.loads(db_get(CandidateEvaluation, id=ev.json()["id"]).criteria_json)[0]["name"] == "Thiết kế"
    # interview locked after an evaluation (S6-06 AC4)
    assert client.post(f"/api/v1/interviews/{iid}/cancel", headers=w["h"]["rec"], json={"reason": "x"}).status_code == 409
    # reopen within 24h, one edit (S6-08 AC3)
    assert client.put(f"/api/v1/referrals/evaluations/{ev.json()['id']}", headers=w["h"]["int"], json={"application_id": app_id, "recommendation": "hold"}).status_code == 409
    assert client.post(f"/api/v1/referrals/evaluations/{ev.json()['id']}/reopen", headers=w["h"]["rec"]).status_code == 200
    assert client.put(f"/api/v1/referrals/evaluations/{ev.json()['id']}", headers=w["h"]["int"], json={"application_id": app_id, "recommendation": "hold"}).json()["recommendation"] == "hold"
    summary = client.get(f"/api/v1/referrals/evaluations/summary?application_id={app_id}", headers=w["h"]["hm"]).json()
    assert summary["criteria_averages"]["Thiết kế"] == 5 and summary["recommendations"]["hold"] == 1
    compare = client.get(f"/api/v1/referrals/evaluations/compare?job_id={w['job']}", headers=w["h"]["hm"]).json()
    assert compare[0]["application_id"] == app_id and compare[0]["total_score"] == 84.0
    mine = client.get("/api/v1/interviews?mine=true", headers=w["h"]["int"]).json()
    assert mine[0]["my_evaluation_submitted"] is True  # S7-01 AC2


def test_hiring_decision_logged(w):
    app_id = apply(w, cv=False).json()["id"]
    client.post("/api/v1/interviews", headers=w["h"]["rec"], json={"application_id": app_id, "scheduled_time": _when(30), "interviewer_ids": [w["ids"]["hm"]]})
    assert client.post(f"/api/v1/applications/{app_id}/hiring-decision", headers=w["h"]["hm"], json={"decision": "reject"}).status_code == 422
    ok = client.post(f"/api/v1/applications/{app_id}/hiring-decision", headers=w["h"]["hm"], json={"decision": "offer", "reason": "Phù hợp"})
    assert ok.status_code == 200
    assert any("Quyết định tuyển" in (l["comment"] or "") for l in client.get(f"/api/v1/applications/{app_id}", headers=w["h"]["hm"]).json()["logs"])


# ------------------------------------------------------------------ offers via e-mail link (S7-04..S7-08)
def test_offer_letter_link_acceptance_and_department_onboarding_template(w):
    client.post("/api/v1/hr/onboarding-templates", headers=w["h"]["hrm"], json={
        "name": "Kỹ thuật", "department_id": w["dept"], "items": [{"task_name": "Cấp laptop", "assigned_role": "IT", "days_before_start": 2},
                                                                   {"task_name": "Cấp tài khoản Git", "assigned_role": "IT"}]})
    app_id = apply(w, cv=False).json()["id"]
    client.post("/api/v1/interviews", headers=w["h"]["rec"], json={"application_id": app_id, "scheduled_time": _when(30), "interviewer_ids": [w["ids"]["int"]]})
    start = (date.today() + timedelta(days=30)).isoformat()
    offer = client.post("/api/v1/offers", headers=w["h"]["rec"], json={"application_id": app_id, "offered_salary": 25_000_000, "allowance": 2_000_000, "start_date": start}).json()
    preview = client.get(f"/api/v1/offers/{offer['id']}/letter-preview", headers=w["h"]["rec"]).json()
    assert "25,000,000" in preview["body_html"]
    client.post(f"/api/v1/offers/{offer['id']}/decision", headers=w["h"]["hrm"], json={"action": "approved"})
    client.post(f"/api/v1/offers/{offer['id']}/send", headers=w["h"]["rec"])
    body = db_get(EmailLog, template_code="OFFER_LETTER", related_id=str(app_id)).body_html
    token = body.split("token=")[1].split('"')[0]
    assert client.get(f"/api/v1/public/offers?token={token}").json()["allowance"] == 2_000_000
    assert client.post("/api/v1/public/offers/respond", json={"token": token, "decision": "accepted"}).json()["status"] == "accepted"
    assert client.post("/api/v1/public/offers/respond", json={"token": token, "decision": "accepted"}).status_code == 404  # single use
    data = client.get(f"/api/v1/offers/{offer['id']}", headers=w["h"]["hrm"]).json()
    assert [i["task_name"] for i in data["onboarding"]] == ["Cấp laptop", "Cấp tài khoản Git"]
    assert db_get(JobApplication, id=app_id).status == "hired"


# ------------------------------------------------------------------ SLA (S5-10) and reports (S8-05..S8-09)
def test_sla_flags_and_reminds_once_per_day(w):
    db = TestingSessionLocal()
    db.add(SlaRule(stage="submitted", max_days=3)); db.commit(); db.close()
    app_id = apply(w, cv=False).json()["id"]
    db = TestingSessionLocal()
    db.query(JobApplication).filter(JobApplication.id == app_id).update({"stage_entered_at": datetime.utcnow() - timedelta(days=5)})
    db.commit()
    assert run_sla_check(db)["reminded"] == 1
    assert run_sla_check(db)["reminded"] == 0
    db.close()
    assert db_get(JobApplication, id=app_id).is_overdue is True
    assert client.get("/api/v1/applications?overdue_only=true", headers=w["h"]["rec"]).json()[0]["id"] == app_id


def test_reports_from_real_data(w):
    ids = [apply(w, email=f"s{i}@example.com", cv=False, source=src).json()["id"] for i, src in enumerate(["LinkedIn", "LinkedIn", "TopCV"])]
    client.put(f"/api/v1/applications/{ids[0]}/status", headers=w["h"]["rec"], json={"status": "screening"})
    funnel = client.get(f"/api/v1/analytics/funnel?job_id={w['job']}", headers=w["h"]["hrm"]).json()
    assert funnel["total"] == 3 and funnel["stages"][1]["count"] == 1 and funnel["biggest_drop"]["from"] == "submitted"
    sources = client.get("/api/v1/analytics/sources", headers=w["h"]["hrm"]).json()
    assert {s["source"]: s["applications"] for s in sources} == {"LinkedIn": 2, "TopCV": 1}
    client.post("/api/v1/analytics/costs", headers=w["h"]["hrm"], json={"requisition_id": w["req"], "category": "posting", "amount": 3_000_000})
    metrics = client.get("/api/v1/analytics/hiring-metrics", headers=w["h"]["hrm"]).json()
    assert metrics["total_hires"] == 0 and metrics["cost_per_hire"] is None
    xlsx = client.get("/api/v1/analytics/sources?export=true", headers=w["h"]["hrm"])
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    trends = client.get("/api/v1/analytics/trends", headers=w["h"]["hrm"]).json()
    assert sum(x["applications"] for x in trends["weekly_applications"]) == 3
    hm = client.get("/api/v1/analytics/hiring-manager", headers=w["h"]["hm"]).json()
    assert hm[0]["code"] == "REQ-F-1" and hm[0]["stage_counts"]["submitted"] == 2
    assert client.get("/api/v1/analytics/funnel", headers=w["h"]["emp"]).status_code == 403


def test_direct_email_is_escaped_scoped_and_queued(w):
    app_id = apply(w, cv=False).json()["id"]
    body = {"subject": "Hẹn trao đổi", "body": "Xin chào <script>alert(1)</script>\nHẹn gặp bạn"}
    assert client.post(f"/api/v1/applications/{app_id}/email", headers=w["h"]["rec2"], json=body).status_code == 403  # not in scope
    ok = client.post(f"/api/v1/applications/{app_id}/email", headers=w["h"]["rec"], json=body)
    assert ok.status_code == 201, ok.text
    log = db_get(EmailLog, id=ok.json()["id"])
    assert log.status == "queued" and log.template_code == "DIRECT_MESSAGE"
    assert "<script>" not in log.body_html and "&lt;script&gt;" in log.body_html and "<br>" in log.body_html
