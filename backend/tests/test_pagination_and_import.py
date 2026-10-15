import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.main import app as fastapi_app
from app.db.session import Base, get_db
from app.db.seed_rbac import ROLES_DATA, PERMISSIONS_DATA, DEFAULT_GRANTED_MATRIX
from app.models.user import User
from app.models.rbac import Role, UserRole
from app.models.audit import Permission, RolePermission
from app.models.import_job import ImportJob
from app.core.security import hash_password
from app.core.pagination import paginate_query
from tests.conftest import session_token

engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


fastapi_app.dependency_overrides[get_db] = override_get_db
client = TestClient(fastapi_app)


@pytest.fixture(autouse=True)
def setup_db():
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    # Seed roles
    roles = {r["code"]: Role(code=r["code"], name=r["name"]) for r in ROLES_DATA}
    db.add_all(roles.values())
    for p in PERMISSIONS_DATA:
        db.add(Permission(code=p["code"], name=p["name"], module=p["module"]))
    for role_code, granted in DEFAULT_GRANTED_MATRIX.items():
        for p in PERMISSIONS_DATA:
            db.add(RolePermission(role_code=role_code, permission_code=p["code"], is_granted=p["code"] in granted))
    db.commit()

    # Seed Admin User
    admin = User(
        email="admin@example.com",
        full_name="Quản Trị Viên",
        hashed_password=hash_password("Password123!"),
        role="admin",
        is_active=True,
        account_status="active",
        security_version=1,
        department="Ban Giám Đốc",
        employee_code="NV-ADM-001",
    )
    db.add(admin)
    db.flush()
    db.add(UserRole(user_id=admin.id, role_id=roles["admin"].id))

    # Seed sample users for pagination testing
    for i in range(1, 15):
        u = User(
            email=f"user_{i:02d}@example.com",
            full_name=f"Nhân Viên {i:02d}",
            hashed_password=hash_password("Password123!"),
            role="employee",
            is_active=True,
            account_status="active",
            security_version=1,
            department="Nhân sự" if i % 2 == 0 else "Kỹ thuật",
            employee_code=f"NV-{i:03d}",
        )
        db.add(u)
        db.flush()
        db.add(UserRole(user_id=u.id, role_id=roles["employee"].id))

    db.commit()
    db.close()
    yield
    Base.metadata.drop_all(bind=engine)


@pytest.fixture
def admin_auth_header():
    """Create session token for seeded admin."""
    db = TestingSessionLocal()
    try:
        admin = db.query(User).filter(User.email == "admin@example.com").first()
        token = session_token(db, admin)
        return {"Authorization": f"Bearer {token}"}
    finally:
        db.close()


def test_paginate_query_utility():
    """Test core paginate_query helper for sorting, counting, and deterministic ordering."""
    db = TestingSessionLocal()
    try:
        query = db.query(User)
        result = paginate_query(
            query=query,
            page=1,
            page_size=5,
            sort_by="email",
            sort_order="asc",
            model=User,
            allowed_sort_columns=["email", "created_at", "full_name"],
        )
        assert "items" in result
        assert "total" in result
        assert result["page"] == 1
        assert result["page_size"] == 5
        assert len(result["items"]) == 5
        assert result["pages"] >= 1
    finally:
        db.close()


def test_list_users_pagination_and_sorting(admin_auth_header):
    """Test GET /api/v1/admin/users pagination, sorting, and department filtering."""
    # 1. Default pagination
    res = client.get("/api/v1/admin/users?page=1&pageSize=5", headers=admin_auth_header)
    assert res.status_code == 200, res.text
    data = res.json()
    assert "items" in data
    assert "total" in data
    assert data["page"] == 1
    assert data["page_size"] == 5
    assert len(data["items"]) == 5

    # 2. Sort by full_name asc
    res_asc = client.get("/api/v1/admin/users?sortBy=full_name&sortOrder=asc&page=1&pageSize=5", headers=admin_auth_header)
    assert res_asc.status_code == 200
    names_asc = [u["full_name"] for u in res_asc.json()["items"]]

    # 3. Sort by full_name desc
    res_desc = client.get("/api/v1/admin/users?sortBy=full_name&sortOrder=desc&page=1&pageSize=5", headers=admin_auth_header)
    assert res_desc.status_code == 200
    names_desc = [u["full_name"] for u in res_desc.json()["items"]]
    assert names_asc != names_desc

    # 4. Filter by department
    res_dept = client.get("/api/v1/admin/users?department=Nhân sự", headers=admin_auth_header)
    assert res_dept.status_code == 200
    for u in res_dept.json()["items"]:
        assert u["department"] == "Nhân sự"


def test_download_import_template(admin_auth_header):
    """Test GET /api/v1/admin/users/import/template produces valid Excel file."""
    res = client.get("/api/v1/admin/users/import/template", headers=admin_auth_header)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["Content-Type"]
    assert "mau_import_nguoi_dung.xlsx" in res.headers["Content-Disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content), data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert "Họ và tên" in headers
    assert "Email" in headers
    assert "Vai trò" in headers
    assert "Phòng ban" in headers


def test_excel_import_preview_and_commit_workflow(admin_auth_header):
    """Test complete S2-01 workflow: preview with error detection -> paginated preview -> batch commit with user_roles."""
    # 1. Create in-memory Excel file with 3 rows: 1 valid, 1 invalid email, 1 duplicate email
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Họ và tên", "Email", "Vai trò", "Số điện thoại", "Phòng ban", "Mã nhân viên", "Mật khẩu ban đầu"])
    ws.append(["Nguyễn Văn Hợp Lệ", "user_valid_import_1@example.com", "recruiter", "0901234567", "Nhân sự", "NV-IMP-001", "Password123!"])
    ws.append(["Lê Văn Lỗi Email", "not-an-email", "employee", "0902345678", "Kỹ thuật", "NV-IMP-002", ""])
    ws.append(["Trần Văn Trùng Admin", "admin@example.com", "employee", "0903456789", "Nhân sự", "NV-IMP-003", ""])

    excel_buf = io.BytesIO()
    wb.save(excel_buf)
    excel_buf.seek(0)

    # 2. Upload to Preview endpoint
    files = {"file": ("test_import.xlsx", excel_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview_res = client.post("/api/v1/admin/users/import/preview", files=files, headers=admin_auth_header)
    assert preview_res.status_code == 200, preview_res.text
    preview_data = preview_res.json()

    assert "job_id" in preview_data
    job_id = preview_data["job_id"]
    assert preview_data["total"] == 3
    assert preview_data["valid_count"] == 1
    assert preview_data["error_count"] == 2

    # Check row-level statuses and error reasons (S2-01 AC2)
    items = preview_data["items"]
    assert len(items) == 3
    assert items[0]["status"] == "VALID"
    assert items[0]["email"] == "user_valid_import_1@example.com"
    assert items[1]["status"] == "ERROR"
    assert "không đúng định dạng" in items[1]["error_reason"]
    assert items[2]["status"] == "DUPLICATE"
    assert "đã tồn tại" in items[2]["error_reason"]

    # 3. Query paginated preview filtering for errors only
    page_res = client.get(f"/api/v1/admin/users/import/preview/{job_id}?status_filter=error&page=1&pageSize=10", headers=admin_auth_header)
    assert page_res.status_code == 200
    error_page = page_res.json()
    assert error_page["total"] == 2
    for item in error_page["items"]:
        assert item["status"] in ("ERROR", "DUPLICATE")

    # 4. Commit import (S2-01 AC3: valid rows imported, errors skipped)
    commit_res = client.post("/api/v1/admin/users/import/commit", json={"job_id": job_id, "skip_errors": True}, headers=admin_auth_header)
    assert commit_res.status_code == 200, commit_res.text
    commit_data = commit_res.json()

    assert commit_data["total"] == 3
    assert commit_data["imported"] == 1
    assert commit_data["failed"] >= 1

    # 5. Verify database persistence and NORMALIZED user_roles assignment (Critical Defect Fix)
    db = TestingSessionLocal()
    try:
        user = db.query(User).filter(User.email == "user_valid_import_1@example.com").first()
        assert user is not None
        assert user.full_name == "Nguyễn Văn Hợp Lệ"
        assert user.department == "Nhân sự"
        assert "recruiter" in user.role

        # Check normalized user_roles table
        roles = db.query(Role).join(UserRole, UserRole.role_id == Role.id).filter(UserRole.user_id == user.id).all()
        role_codes = {r.code for r in roles}
        assert "recruiter" in role_codes
    finally:
        db.close()

    # 6. Idempotent commit check
    repeat_commit = client.post("/api/v1/admin/users/import/commit", json={"job_id": job_id, "skip_errors": True}, headers=admin_auth_header)
    assert repeat_commit.status_code == 200
    assert repeat_commit.json()["imported"] == 1


def test_import_missing_columns_rejected(admin_auth_header):
    """Test S2-01 header validation: file missing required columns is rejected with 400."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tên", "Điện thoại"])  # Missing Email and Vai trò
    ws.append(["Ai Đó", "0901234567"])

    excel_buf = io.BytesIO()
    wb.save(excel_buf)
    excel_buf.seek(0)

    files = {"file": ("missing_cols.xlsx", excel_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post("/api/v1/admin/users/import/preview", files=files, headers=admin_auth_header)
    assert res.status_code == 400
    assert "Thiếu các cột bắt buộc" in res.json()["detail"]
